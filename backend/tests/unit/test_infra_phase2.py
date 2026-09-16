# ============================================================
# OPSYN INFRASTRUCTURE — Phase 2 Unit Tests
# tests/unit/test_infra_phase2.py
#
# Covers all validation gates from the implementation plan:
#  Gate 1: validate_coordinates out-of-range → error
#  Gate 2: validate_coordinates valid → no error
#  Gate 3: validate_row_splitter with splitter_level='first-level' → exact error
#  Gate 4: run_full_validation 10-row file, 2 errors → error_rows=2, valid_rows=8
#  Gate 5: commit_session inserts valid rows, marks staging committed=True
#  Gate 6: commit_session DB error → rollback all, session=rolled_back
#  Gate 7: generate_error_report_xlsx produces valid .xlsx with Errors column
#
# Run: pytest tests/unit/test_infra_phase2.py -v
# ============================================================

import io
import uuid
import asyncio
from unittest.mock import AsyncMock, MagicMock, patch, call
import pytest
import openpyxl

from app.modules.infrastructure.gis_utils import (
    validate_coordinates, make_point, warn_zero_island,
)
from app.modules.infrastructure.upload_service import (
    validate_file, validate_headers, validate_row_splitter,
    validate_row_cabinet, validate_row_olt,
    run_full_validation, commit_session, generate_error_report_xlsx,
    REQUIRED_COLUMNS,
)
from app.modules.infrastructure.schemas import ValidationSummary, CommitResult
from app.modules.infrastructure.audit_service import write_audit_entry


# ══════════════════════════════════════════════════════════════
#  Helpers
# ══════════════════════════════════════════════════════════════

def _make_splitter_xlsx(rows: list[dict]) -> bytes:
    """Build a minimal splitter .xlsx in memory."""
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["box_id", "input", "output", "splitter_level",
               "splitter_type", "number_customer", "longitude", "latitude"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _valid_splitter_row(box_id="SPL-001", level="First-Level") -> dict:
    return {
        "box_id": box_id, "input": 1, "output": 8,
        "splitter_level": level, "splitter_type": "PCC",
        "number_customer": 4, "longitude": 3.3792, "latitude": 6.5244,
    }


# ══════════════════════════════════════════════════════════════
#  Gate 1 & 2 — validate_coordinates
# ══════════════════════════════════════════════════════════════

class TestValidateCoordinates:

    def test_longitude_out_of_range(self):
        """Gate 1: (200, 45) must return invalid with an error."""
        valid, errors = validate_coordinates(200, 45)
        assert not valid
        assert len(errors) > 0
        assert any("180" in e for e in errors), f"Expected range error, got: {errors}"

    def test_latitude_out_of_range(self):
        valid, errors = validate_coordinates(10, 95)
        assert not valid
        assert any("90" in e for e in errors)

    def test_valid_coordinates(self):
        """Gate 2: (45.3, 6.2) must return valid with no errors."""
        valid, errors = validate_coordinates(45.3, 6.2)
        assert valid
        assert errors == []

    def test_boundary_values_accepted(self):
        for lng, lat in [(-180, 0), (180, 0), (0, -90), (0, 90)]:
            valid, errors = validate_coordinates(lng, lat)
            assert valid, f"Boundary ({lng},{lat}) should be valid; errors: {errors}"

    def test_missing_values(self):
        valid, errors = validate_coordinates(None, None)
        assert not valid
        assert any("required" in e.lower() for e in errors)

    def test_non_numeric_longitude(self):
        valid, errors = validate_coordinates("abc", 6.0)
        assert not valid
        assert any("valid number" in e.lower() for e in errors)

    def test_zero_island_warning(self):
        assert warn_zero_island(0.0, 0.0) is True
        assert warn_zero_island(1.0, 2.0) is False


# ══════════════════════════════════════════════════════════════
#  Gate 3 — validate_row_splitter enum case-sensitivity
# ══════════════════════════════════════════════════════════════

class TestValidateRowSplitter:

    def test_lowercase_level_returns_exact_error(self):
        """Gate 3: 'first-level' must produce exact error message."""
        row = _valid_splitter_row(level="first-level")
        result = validate_row_splitter(row)
        assert not result.valid
        error_messages = [e["error"] for e in result.errors]
        assert any(
            "splitter_level must be exactly 'First-Level' or 'Second-Level'" == msg
            for msg in error_messages
        ), f"Exact error not found. Got: {error_messages}"

    def test_allcaps_level_invalid(self):
        row = _valid_splitter_row(level="FIRST-LEVEL")
        result = validate_row_splitter(row)
        assert not result.valid

    def test_spaced_level_invalid(self):
        row = _valid_splitter_row(level="First Level")  # space instead of hyphen
        result = validate_row_splitter(row)
        assert not result.valid

    def test_valid_splitter_row(self):
        result = validate_row_splitter(_valid_splitter_row())
        assert result.valid
        assert result.errors == []

    def test_second_level_valid(self):
        result = validate_row_splitter(_valid_splitter_row(level="Second-Level"))
        assert result.valid

    def test_invalid_splitter_type(self):
        row = {**_valid_splitter_row(), "splitter_type": "pcc"}
        result = validate_row_splitter(row)
        assert not result.valid
        assert any("splitter_type must be exactly 'PCC' or 'Legacy'" in e["error"]
                   for e in result.errors)

    def test_missing_required_field(self):
        row = {**_valid_splitter_row(), "box_id": None}
        result = validate_row_splitter(row)
        assert not result.valid

    def test_input_below_minimum(self):
        row = {**_valid_splitter_row(), "input": 0}
        result = validate_row_splitter(row)
        assert not result.valid

    def test_invalid_longitude(self):
        row = {**_valid_splitter_row(), "longitude": 200}
        result = validate_row_splitter(row)
        assert not result.valid

    def test_zero_island_warning(self):
        row = {**_valid_splitter_row(), "longitude": 0, "latitude": 0}
        result = validate_row_splitter(row)
        # (0,0) is a WARNING, not an error — row should still be valid
        assert result.valid
        assert any("0, 0" in w or "0,0" in w for w in result.warnings)

    def test_collects_all_errors(self):
        """Validation must collect all errors in a single pass, not stop early."""
        row = {
            "box_id": None, "input": 0, "output": -1,
            "splitter_level": "bad", "splitter_type": "bad",
            "longitude": 999, "latitude": 999,
        }
        result = validate_row_splitter(row)
        assert not result.valid
        assert len(result.errors) >= 4, \
            f"Expected at least 4 errors collected, got {len(result.errors)}"


class TestValidateRowCabinet:
    def test_valid_cabinet(self):
        row = {"cabinet_id": "CAB-001", "capacity": 48, "number_tray": 4,
               "longitude": 3.38, "latitude": 6.52}
        assert validate_row_cabinet(row).valid

    def test_zero_capacity_invalid(self):
        row = {"cabinet_id": "CAB-001", "capacity": 0, "number_tray": 4,
               "longitude": 3.38, "latitude": 6.52}
        assert not validate_row_cabinet(row).valid


class TestValidateRowOLT:
    def test_valid_olt(self):
        row = {"name": "OLT-01", "location": "Lagos POP",
               "longitude": 3.38, "latitude": 6.52}
        assert validate_row_olt(row).valid

    def test_missing_name(self):
        row = {"name": "", "location": "Lagos POP",
               "longitude": 3.38, "latitude": 6.52}
        assert not validate_row_olt(row).valid

    def test_invalid_number_of_odf(self):
        row = {"name": "OLT-01", "location": "Lagos",
               "number_of_odf": "abc",
               "longitude": 3.38, "latitude": 6.52}
        assert not validate_row_olt(row).valid


# ══════════════════════════════════════════════════════════════
#  Gate 4 — run_full_validation (10 rows, 2 errors)
# ══════════════════════════════════════════════════════════════

class TestRunFullValidation:

    @pytest.mark.asyncio
    async def test_ten_rows_two_errors(self):
        """Gate 4: 10-row file with 2 intentional errors → error_rows=2, valid_rows=8."""
        session_id = uuid.uuid4()
        tenant_id  = uuid.uuid4()

        # Build 8 valid rows + 2 with invalid splitter_level
        rows = [_valid_splitter_row(box_id=f"SPL-{i:03d}") for i in range(8)]
        rows.append({**_valid_splitter_row(box_id="SPL-ERR-001"), "splitter_level": "first-level"})
        rows.append({**_valid_splitter_row(box_id="SPL-ERR-002"), "splitter_level": "FIRST-LEVEL"})
        file_bytes = _make_splitter_xlsx(rows)

        # Mock InfraUploadSession
        mock_session = MagicMock()
        mock_session.id     = session_id
        mock_session.status = "pending"

        # Mock DB
        db = AsyncMock()
        db.get = AsyncMock(return_value=mock_session)
        db.flush  = AsyncMock()
        db.add_all = MagicMock()

        # Duplicate checks all return False (no existing records)
        async def fake_execute(stmt, *args, **kwargs):
            mock_result = MagicMock()
            mock_result.scalar.return_value = None
            return mock_result

        db.execute = AsyncMock(side_effect=fake_execute)

        summary = await run_full_validation(
            db, session_id, "splitter", file_bytes, tenant_id,
            filename="test.xlsx",
        )

        assert summary.total_rows    == 10
        assert summary.valid_rows    == 8
        assert summary.error_rows    == 2
        assert summary.duplicate_rows == 0
        assert summary.status        == "validated"

    @pytest.mark.asyncio
    async def test_invalid_file_extension_returns_failed(self):
        db = AsyncMock()
        mock_session = MagicMock()
        mock_session.status = "pending"
        db.get = AsyncMock(return_value=mock_session)
        db.flush = AsyncMock()

        summary = await run_full_validation(
            db, uuid.uuid4(), "splitter", b"fake data",
            uuid.uuid4(), filename="data.csv",
        )
        assert summary.status == "failed"

    @pytest.mark.asyncio
    async def test_missing_header_returns_failed(self):
        """File with missing required column → status=failed."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["box_id", "input"])  # missing many required columns
        ws.append(["SPL-001", 1])
        buf = io.BytesIO()
        wb.save(buf)
        file_bytes = buf.getvalue()

        mock_session = MagicMock()
        mock_session.status = "pending"
        db = AsyncMock()
        db.get   = AsyncMock(return_value=mock_session)
        db.flush = AsyncMock()

        summary = await run_full_validation(
            db, uuid.uuid4(), "splitter", file_bytes, uuid.uuid4(),
            filename="test.xlsx",
        )
        assert summary.status == "failed"


# ══════════════════════════════════════════════════════════════
#  Gate 5 — commit_session inserts rows, marks committed=True
# ══════════════════════════════════════════════════════════════

class TestCommitSession:

    def _make_staging_rows(self, count: int, asset_type: str = "splitter"):
        rows = []
        for i in range(count):
            r = MagicMock()
            r.row_data = {
                "box_id": f"SPL-{i:03d}", "input": "1", "output": "8",
                "splitter_level": "First-Level", "splitter_type": "PCC",
                "number_customer": "4",
                "longitude": "3.3792", "latitude": "6.5244",
            }
            r.is_valid     = True
            r.is_duplicate = False
            r.committed    = False
            rows.append(r)
        return rows

    @pytest.mark.asyncio
    async def test_commit_sets_committed_true(self):
        """Gate 5: each staging row's committed flag is set to True."""
        session_id = uuid.uuid4()
        tenant_id  = uuid.uuid4()
        approver_id = uuid.uuid4()

        mock_session = MagicMock()
        mock_session.asset_type   = "splitter"
        mock_session.status       = "validated"
        mock_session.committed_rows = None

        staging_rows = self._make_staging_rows(3)

        db = AsyncMock()

        async def _db_get(model_cls, pk):
            if hasattr(mock_session, "asset_type"):
                return mock_session
            return None

        db.get   = AsyncMock(side_effect=_db_get)
        db.flush = AsyncMock()
        db.add   = MagicMock()

        # begin_nested returns an async context manager
        savepoint_cm = AsyncMock()
        savepoint_cm.__aenter__ = AsyncMock(return_value=savepoint_cm)
        savepoint_cm.__aexit__  = AsyncMock(return_value=False)
        db.begin_nested = AsyncMock(return_value=savepoint_cm)

        # execute: first call returns the actor user; second returns staging rows
        actor_user_mock = MagicMock()
        actor_user_mock.id       = approver_id
        actor_user_mock.username = "approver"
        actor_user_mock.staff_profile = None

        staging_result = MagicMock()
        staging_result.scalars.return_value.all.return_value = staging_rows

        execute_calls = [0]
        async def _execute(stmt, *a, **kw):
            execute_calls[0] += 1
            if execute_calls[0] == 1:
                # User query
                r = MagicMock()
                r.scalar_one_or_none.return_value = actor_user_mock
                return r
            # Staging rows query
            return staging_result

        db.execute = AsyncMock(side_effect=_execute)

        with patch(
            "app.modules.infrastructure.upload_service.write_audit_entry",
            new=AsyncMock(),
        ):
            result = await commit_session(db, session_id, tenant_id, approver_id)

        assert result.status        == "committed"
        assert result.committed_rows == 3
        for row in staging_rows:
            assert row.committed is True

    @pytest.mark.asyncio
    async def test_commit_rollback_on_db_error(self):
        """Gate 6: DB error inside savepoint → status=rolled_back, 0 rows committed."""
        session_id  = uuid.uuid4()
        tenant_id   = uuid.uuid4()
        approver_id = uuid.uuid4()

        mock_session = MagicMock()
        mock_session.asset_type = "splitter"
        mock_session.status     = "validated"

        staging_rows = self._make_staging_rows(3)

        db = AsyncMock()
        db.get   = AsyncMock(return_value=mock_session)
        db.flush = AsyncMock()
        db.add   = MagicMock()

        # begin_nested raises on __aenter__ to simulate DB error inside savepoint
        class _FailingSavepoint:
            async def __aenter__(self):
                raise RuntimeError("Simulated DB error on row 3")
            async def __aexit__(self, *a):
                return False

        db.begin_nested = AsyncMock(return_value=_FailingSavepoint())

        actor_mock = MagicMock()
        actor_mock.id = approver_id
        actor_mock.username = "approver"
        actor_mock.staff_profile = None

        staging_result = MagicMock()
        staging_result.scalars.return_value.all.return_value = staging_rows

        call_count = [0]
        async def _execute(stmt, *a, **kw):
            call_count[0] += 1
            if call_count[0] == 1:
                r = MagicMock()
                r.scalar_one_or_none.return_value = actor_mock
                return r
            return staging_result

        db.execute = AsyncMock(side_effect=_execute)

        result = await commit_session(db, session_id, tenant_id, approver_id)

        assert result.status        == "rolled_back"
        assert result.committed_rows == 0
        assert result.error is not None
        assert mock_session.status  == "rolled_back"


# ══════════════════════════════════════════════════════════════
#  Gate 7 — generate_error_report_xlsx
# ══════════════════════════════════════════════════════════════

class TestGenerateErrorReportXlsx:

    @pytest.mark.asyncio
    async def test_produces_valid_xlsx_with_errors_column(self):
        """Gate 7: output must be a valid .xlsx with an 'Errors' column."""
        session_id = uuid.uuid4()

        mock_session = MagicMock()
        mock_session.asset_type = "splitter"

        row1 = MagicMock()
        row1.row_number = 3
        row1.row_data   = {"box_id": "SPL-ERR-001", "splitter_level": "first-level",
                           "longitude": "3.38", "latitude": "6.52"}
        row1.errors     = [{"field": "splitter_level",
                            "error": "splitter_level must be exactly 'First-Level' or 'Second-Level'"}]

        row2 = MagicMock()
        row2.row_number = 7
        row2.row_data   = {"box_id": "SPL-ERR-002", "longitude": "abc"}
        row2.errors     = [{"field": "longitude", "error": "Longitude must be a valid number"}]

        db = AsyncMock()

        get_calls = [0]
        async def _get(cls, pk):
            get_calls[0] += 1
            return mock_session

        db.get = AsyncMock(side_effect=_get)

        invalid_result = MagicMock()
        invalid_result.scalars.return_value.all.return_value = [row1, row2]
        db.execute = AsyncMock(return_value=invalid_result)

        buf = await generate_error_report_xlsx(db, session_id)

        # Verify it's a valid xlsx
        wb = openpyxl.load_workbook(buf)
        ws = wb.active

        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        assert "Errors" in headers, f"'Errors' column missing. Headers: {headers}"
        assert "Row #" in headers

        # Should have 2 data rows + 1 header
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(data_rows) == 2

        # Errors cell should be non-empty for both rows
        errors_col_idx = headers.index("Errors")
        for dr in data_rows:
            assert dr[errors_col_idx] is not None and dr[errors_col_idx] != "", \
                f"Errors column empty for row: {dr}"


# ══════════════════════════════════════════════════════════════
#  validate_file tests
# ══════════════════════════════════════════════════════════════

class TestValidateFile:

    def test_csv_rejected(self):
        result = validate_file(b"col1,col2\nval1,val2", "data.csv")
        assert not result.valid
        assert any(".xlsx" in e for e in result.errors)

    def test_empty_file_rejected(self):
        result = validate_file(b"", "data.xlsx")
        assert not result.valid

    def test_oversized_file_rejected(self):
        # Start with xlsx magic bytes so the MIME check passes,
        # then pad to 11 MB to trigger the size check.
        _XLSX_MAGIC = b'PK\x03\x04'
        big = _XLSX_MAGIC + b"x" * (11 * 1024 * 1024 - len(_XLSX_MAGIC))
        result = validate_file(big, "data.xlsx")
        assert not result.valid
        assert any("MB" in e for e in result.errors)

    def test_valid_xlsx_accepted(self):
        buf = _make_splitter_xlsx([_valid_splitter_row()])
        result = validate_file(buf, "upload.xlsx")
        assert result.valid

    def test_corrupted_xlsx_rejected(self):
        result = validate_file(b"this is not an xlsx", "data.xlsx")
        assert not result.valid


# ══════════════════════════════════════════════════════════════
#  validate_headers tests
# ══════════════════════════════════════════════════════════════

class TestValidateHeaders:

    def _make_ws(self, headers: list[str]):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        return ws

    def test_all_required_headers_present(self):
        ws = self._make_ws(["box_id", "input", "output", "splitter_level",
                            "splitter_type", "longitude", "latitude"])
        result, header_map = validate_headers(ws, "splitter")
        assert result.valid
        assert "box_id" in header_map

    def test_case_insensitive_match(self):
        ws = self._make_ws(["BOX_ID", "INPUT", "OUTPUT", "SPLITTER_LEVEL",
                            "SPLITTER_TYPE", "LONGITUDE", "LATITUDE"])
        result, header_map = validate_headers(ws, "splitter")
        assert result.valid

    def test_missing_column(self):
        ws = self._make_ws(["box_id", "input"])  # missing many
        result, _ = validate_headers(ws, "splitter")
        assert not result.valid
        assert "output" in result.missing_columns or len(result.missing_columns) > 0

    def test_extra_columns_ignored(self):
        ws = self._make_ws(["box_id", "input", "output", "splitter_level",
                            "splitter_type", "longitude", "latitude", "extra_col"])
        result, _ = validate_headers(ws, "splitter")
        assert result.valid


# ══════════════════════════════════════════════════════════════
#  audit_service tests
# ══════════════════════════════════════════════════════════════

class TestAuditService:

    @pytest.mark.asyncio
    async def test_actor_label_from_staff_profile(self):
        sp = MagicMock()
        sp.full_name = "David Onoja"
        actor = MagicMock()
        actor.id = uuid.uuid4()
        actor.username = "david.onoja"
        actor.staff_profile = sp

        db = AsyncMock()
        db.add   = MagicMock()
        db.flush = AsyncMock()

        entry = await write_audit_entry(
            db,
            asset_type="splitter", asset_id=uuid.uuid4(), asset_key="SPL-001",
            action="upload", actor_user=actor,
            tenant_id=uuid.uuid4(),
        )
        assert entry.actor_label == "David Onoja"

    @pytest.mark.asyncio
    async def test_actor_label_fallback_to_username(self):
        actor = MagicMock()
        actor.id = uuid.uuid4()
        actor.username = "jane.doe"
        actor.staff_profile = None

        db = AsyncMock()
        db.add   = MagicMock()
        db.flush = AsyncMock()

        entry = await write_audit_entry(
            db,
            asset_type="cabinet", asset_id=uuid.uuid4(), asset_key="CAB-001",
            action="soft_delete", actor_user=actor,
            tenant_id=uuid.uuid4(),
        )
        assert entry.actor_label == "jane.doe"

    @pytest.mark.asyncio
    async def test_actor_label_system_when_no_user(self):
        db = AsyncMock()
        db.add   = MagicMock()
        db.flush = AsyncMock()

        entry = await write_audit_entry(
            db,
            asset_type="olt", asset_id=uuid.uuid4(), asset_key="OLT-01",
            action="upload", actor_user=None,
            tenant_id=uuid.uuid4(),
        )
        assert entry.actor_label == "System"
