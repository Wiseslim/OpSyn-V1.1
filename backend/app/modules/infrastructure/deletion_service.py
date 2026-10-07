# ============================================================
# OPSYN INFRASTRUCTURE DELETION SERVICE
# app/modules/infrastructure/deletion_service.py
#
# Batch deletion workflow:
#   validate_deletion_file  → parse and validate the .xlsx
#   match_deletion_keys_*   → find which rows match live records
#   execute_deletion        → soft-delete all matched rows atomically
# ============================================================

from __future__ import annotations

import io
import uuid
import datetime
from collections import deque
from typing import Any

import openpyxl
from sqlalchemy import select, update
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.infrastructure.models import (
    InfraDeletionSession, SplitterBox, Cabinet, OLT,
)
from app.modules.staff.models import User
from app.modules.infrastructure.schemas import (
    DeletionValidationResult, MatchResult, DeletionResult,
)
from app.modules.infrastructure.audit_service import write_audit_entry

# xlsx is a ZIP container; its first 4 bytes are always PK\x03\x04
_XLSX_MAGIC = b'PK\x03\x04'

# Guard against a corrupt parent_splitter_id cycle turning the walk below into
# an infinite loop. Real FTTH hierarchies are 2-3 deep; this is far above any
# legitimate tree while still terminating on a cycle.
_MAX_CASCADE_DEPTH = 32

# Required columns in deletion files (lowercased for matching)
DELETION_REQUIRED_COLUMNS: dict[str, list[str]] = {
    "splitter": ["box_id", "splitter_level"],
    "cabinet":  ["cabinet_id", "capacity"],
    "olt":      ["name"],
}


# ── File + header parsing ─────────────────────────────────────

async def validate_deletion_file(
    asset_type: str,
    file_bytes: bytes,
    filename: str = "deletion.xlsx",
) -> DeletionValidationResult:
    """Parse .xlsx deletion file and extract deletion keys.

    Returns DeletionValidationResult with the list of parsed key dicts.
    """
    if not filename.lower().endswith(".xlsx"):
        return DeletionValidationResult(
            valid=False, asset_type=asset_type, total_rows=0,
            errors=["Invalid file type — only .xlsx files are accepted"],
        )

    if len(file_bytes) == 0:
        return DeletionValidationResult(
            valid=False, asset_type=asset_type, total_rows=0,
            errors=["Uploaded file is empty"],
        )

    # MIME-type check: xlsx must have ZIP magic bytes PK\x03\x04
    if file_bytes[:4] != _XLSX_MAGIC:
        return DeletionValidationResult(
            valid=False, asset_type=asset_type, total_rows=0,
            errors=["File content does not match .xlsx format — file must be a valid Excel workbook"],
        )

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        return DeletionValidationResult(
            valid=False, asset_type=asset_type, total_rows=0,
            errors=[f"File could not be parsed: {exc}"],
        )

    ws = wb.active
    required_cols = DELETION_REQUIRED_COLUMNS.get(asset_type, [])

    # Parse header row
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        return DeletionValidationResult(
            valid=False, asset_type=asset_type, total_rows=0,
            errors=["File has no header row"],
        )

    header_map: dict[str, int] = {
        str(c).strip().lower(): i
        for i, c in enumerate(header_row)
        if c is not None
    }
    missing = [col for col in required_cols if col not in header_map]
    if missing:
        return DeletionValidationResult(
            valid=False, asset_type=asset_type, total_rows=0,
            errors=[f"Missing required columns: {', '.join(missing)}"],
        )

    keys: list[dict[str, Any]] = []
    for raw_row in ws.iter_rows(min_row=2, values_only=True):
        if all(cell is None for cell in raw_row):
            continue
        row_dict = {col: raw_row[idx] for col, idx in header_map.items()
                    if idx < len(raw_row)}
        key: dict[str, Any] = {}
        for col in required_cols:
            val = row_dict.get(col)
            key[col] = str(val).strip() if val is not None else ""
        keys.append(key)

    wb.close()
    return DeletionValidationResult(
        valid=True,
        asset_type=asset_type,
        total_rows=len(keys),
        errors=[],
        keys=keys,
    )


# ── Match deletion keys against live records ──────────────────

async def match_deletion_keys_splitter(
    db: AsyncSession,
    keys: list[dict[str, Any]],
    tenant_id: uuid.UUID,
) -> list[MatchResult]:
    """Match (box_id, splitter_level) pairs against live splitter_boxes."""
    results: list[MatchResult] = []
    for key in keys:
        box_id = key.get("box_id", "")
        level  = key.get("splitter_level", "")
        row = (await db.execute(
            select(SplitterBox).where(
                SplitterBox.box_id == box_id,
                SplitterBox.splitter_level == level,
                SplitterBox.tenant_id == tenant_id,
                SplitterBox.is_deleted.is_(False),
            ).limit(1)
        )).scalars().first()

        if row:
            results.append(MatchResult(
                key=key, status="matched", asset_id=row.id
            ))
        else:
            results.append(MatchResult(
                key=key, status="unmatched",
                reason="No active record found with this box_id and splitter_level",
            ))
    return results


async def match_deletion_keys_cabinet(
    db: AsyncSession,
    keys: list[dict[str, Any]],
    tenant_id: uuid.UUID,
) -> list[MatchResult]:
    """Match (cabinet_id, capacity) pairs against live cabinets."""
    results: list[MatchResult] = []
    for key in keys:
        cab_id   = key.get("cabinet_id", "")
        capacity_raw = key.get("capacity", "")
        try:
            capacity = int(float(str(capacity_raw)))
        except (TypeError, ValueError):
            results.append(MatchResult(
                key=key, status="unmatched",
                reason=f"Invalid capacity value '{capacity_raw}'",
            ))
            continue

        row = (await db.execute(
            select(Cabinet).where(
                Cabinet.cabinet_id == cab_id,
                Cabinet.capacity == capacity,
                Cabinet.tenant_id == tenant_id,
                Cabinet.is_deleted.is_(False),
            ).limit(1)
        )).scalars().first()

        if row:
            results.append(MatchResult(key=key, status="matched", asset_id=row.id))
        else:
            results.append(MatchResult(
                key=key, status="unmatched",
                reason="No active cabinet found with this cabinet_id and capacity",
            ))
    return results


async def match_deletion_keys_olt(
    db: AsyncSession,
    keys: list[dict[str, Any]],
    tenant_id: uuid.UUID,
) -> list[MatchResult]:
    """Match OLT name against live olts."""
    results: list[MatchResult] = []
    for key in keys:
        name = key.get("name", "")
        row = (await db.execute(
            select(OLT).where(
                OLT.name == name,
                OLT.tenant_id == tenant_id,
                OLT.is_deleted.is_(False),
            ).limit(1)
        )).scalars().first()

        if row:
            results.append(MatchResult(key=key, status="matched", asset_id=row.id))
        else:
            results.append(MatchResult(
                key=key, status="unmatched",
                reason="No active OLT found with this name",
            ))
    return results


# ── Cascade targets ───────────────────────────────────────────
#
# splitter_boxes.cabinet_id and splitter_boxes.parent_splitter_id are both
# ON DELETE SET NULL, so a hard delete of a cabinet or a First-Level splitter
# would silently orphan its children: they would stay live and active while
# pointing at nothing. Soft delete has the same failure in application form,
# because it never reaches the FK at all. So the deletion walk has to carry
# the children with it.


async def _splitter_descendants(
    db: AsyncSession,
    roots: list[uuid.UUID],
    tenant_id: uuid.UUID,
) -> list[uuid.UUID]:
    """Every active splitter below `roots`, breadth-first, cycle-safe.

    Two separate sets, and the distinction matters: `seen` is every child
    discovered (used to report each box once), `expanded` is every node whose
    own children have already been fetched (used to terminate). Collapsing
    them into one set makes the second level of the walk an immediate no-op,
    because the frontier is by then entirely made of "seen" nodes.
    """
    found: list[uuid.UUID] = []
    seen: set[uuid.UUID] = set()
    expanded: set[uuid.UUID] = set()
    frontier: deque[uuid.UUID] = deque(roots)

    for _ in range(_MAX_CASCADE_DEPTH):
        batch = [i for i in frontier if i not in expanded]
        if not batch:
            break
        expanded.update(batch)

        rows = (await db.execute(
            select(SplitterBox.id).where(
                SplitterBox.parent_splitter_id.in_(batch),
                SplitterBox.tenant_id == tenant_id,
                SplitterBox.is_deleted.is_(False),
            )
        )).scalars().all()

        for row_id in rows:
            if row_id not in seen:
                seen.add(row_id)
                found.append(row_id)

        frontier = deque(rows)

    return found


async def collect_cascade_targets(
    db: AsyncSession,
    asset_type: str,
    matched_ids: list[uuid.UUID],
    tenant_id: uuid.UUID,
) -> tuple[list[uuid.UUID], list[uuid.UUID]]:
    """Split the deletion set into explicitly requested and cascaded ids.

    Returns (explicit, cascaded). OLTs have no children and always return an
    empty cascaded list.
    """
    if not matched_ids:
        return [], []

    cascaded: list[uuid.UUID] = []

    if asset_type == "cabinet":
        # Everything mounted in the cabinet, then everything below those.
        direct = (await db.execute(
            select(SplitterBox.id).where(
                SplitterBox.cabinet_id.in_(matched_ids),
                SplitterBox.tenant_id == tenant_id,
                SplitterBox.is_deleted.is_(False),
            )
        )).scalars().all()
        cascaded.extend(direct)
        cascaded.extend(await _splitter_descendants(db, list(direct), tenant_id))

    elif asset_type == "splitter":
        cascaded.extend(await _splitter_descendants(db, matched_ids, tenant_id))

    # De-duplicate while keeping order, and drop anything already requested.
    explicit_set = set(matched_ids)
    seen: set[uuid.UUID] = set()
    unique_cascaded = [
        i for i in cascaded
        if i not in explicit_set and not (i in seen or seen.add(i))
    ]
    return list(matched_ids), unique_cascaded


# ── Execute deletion ──────────────────────────────────────────

_ORM_MAP = {"splitter": SplitterBox, "cabinet": Cabinet, "olt": OLT}


async def execute_deletion(
    db: AsyncSession,
    session_id: uuid.UUID,
    tenant_id: uuid.UUID,
    approved_by_user_id: uuid.UUID,
) -> DeletionResult:
    """Soft-delete all matched assets plus their dependants, atomically.

    Deleting a cabinet also deletes the splitters mounted in it and their
    descendants; deleting a splitter also deletes its descendants. Children
    the approver never named are still reported and audited, so the count in
    the UI is the count that was actually deleted.

    Updates session to executed or failed. Writes audit log for every deleted asset.
    """
    session_obj = await db.get(InfraDeletionSession, session_id)
    if session_obj is None:
        raise ValueError(f"Deletion session {session_id} not found")

    asset_type = session_obj.asset_type
    match_details: list[dict] = session_obj.match_details or []
    matched_ids: list[uuid.UUID] = []

    for detail in match_details:
        if detail.get("status") == "matched" and detail.get("asset_id"):
            matched_ids.append(uuid.UUID(str(detail["asset_id"])))

    # Fetch approver with staff_profile for actor_label
    actor_user = (await db.execute(
        select(User)
        .where(User.id == approved_by_user_id)
        .options(selectinload(User.staff_profile))
    )).scalar_one_or_none()

    orm_cls = _ORM_MAP.get(asset_type)
    if orm_cls is None:
        session_obj.status = "failed"
        await db.flush()
        return DeletionResult(
            session_id=session_id, deleted_rows=0, status="failed",
            error=f"Unsupported asset_type '{asset_type}'",
        )

    explicit_ids, cascaded_ids = await collect_cascade_targets(
        db, asset_type, matched_ids, tenant_id
    )

    deleted_count = 0
    cascaded_count = 0
    try:
        async with await db.begin_nested():  # type: ignore[attr-defined]
            for asset_id in explicit_ids:
                asset_obj = await db.get(orm_cls, asset_id)
                if asset_obj is None or asset_obj.is_deleted:
                    continue
                if asset_obj.tenant_id != tenant_id:
                    continue

                # Capture snapshot before delete
                old_snapshot = _asset_snapshot(asset_type, asset_obj)
                asset_key    = _asset_key_from_obj(asset_type, asset_obj)

                asset_obj.is_deleted = True
                await db.flush()

                await write_audit_entry(
                    db,
                    asset_type=asset_type,
                    asset_id=asset_id,
                    asset_key=asset_key,
                    action="soft_delete",
                    session_id=session_id,
                    actor_user=actor_user,
                    old_data=old_snapshot,
                    new_data=None,
                    tenant_id=tenant_id,
                )
                deleted_count += 1

            # Dependants the approver never named. Set in one statement so a
            # failure part-way through cannot leave half a subtree deleted.
            if cascaded_ids:
                await db.execute(
                    update(SplitterBox)
                    .where(
                        SplitterBox.id.in_(cascaded_ids),
                        SplitterBox.tenant_id == tenant_id,
                        SplitterBox.is_deleted.is_(False),
                    )
                    .values(is_deleted=True)
                )
                cascaded_count = len(cascaded_ids)

                # One audit row per cascaded splitter: the audit log is the
                # only record of why a box the user never named disappeared.
                cascade_rows = (await db.execute(
                    select(
                        SplitterBox.id,
                        SplitterBox.box_id,
                        SplitterBox.splitter_level,
                        SplitterBox.splitter_type,
                    ).where(SplitterBox.id.in_(cascaded_ids))
                )).all()
                for row in cascade_rows:
                    await write_audit_entry(
                        db,
                        asset_type="splitter",
                        asset_id=row.id,
                        asset_key=row.box_id,
                        action="soft_delete_cascade",
                        session_id=session_id,
                        actor_user=actor_user,
                        old_data={
                            "box_id": row.box_id,
                            "splitter_level": row.splitter_level,
                            "splitter_type": row.splitter_type,
                        },
                        new_data=None,
                        tenant_id=tenant_id,
                    )

            # Dual-write (Network GIS Phase 2): retire the net_assets mirrors
            # of what was just deleted, in the same savepoint.
            if deleted_count or cascaded_count:
                from app.modules.network.backfill import mirror_uploaded
                await mirror_uploaded(db, tenant_id)

    except Exception as exc:
        session_obj.status = "failed"
        await db.flush()
        return DeletionResult(
            session_id=session_id, deleted_rows=0, status="failed", error=str(exc)
        )

    # Success
    session_obj.status       = "executed"
    session_obj.deleted_rows = deleted_count + cascaded_count
    session_obj.approved_by  = approved_by_user_id
    session_obj.approved_at  = datetime.datetime.now(datetime.timezone.utc)
    session_obj.executed_at  = datetime.datetime.now(datetime.timezone.utc)
    await db.flush()

    return DeletionResult(
        session_id=session_id,
        deleted_rows=deleted_count,
        cascaded_rows=cascaded_count,
        status="executed",
    )


# ── Helpers ───────────────────────────────────────────────────

def _asset_key_from_obj(asset_type: str, obj) -> str:
    if asset_type == "splitter":
        return getattr(obj, "box_id", "")
    if asset_type == "cabinet":
        return getattr(obj, "cabinet_id", "")
    if asset_type == "olt":
        return getattr(obj, "name", "")
    return ""


def _asset_snapshot(asset_type: str, obj) -> dict[str, Any]:
    """Serialise the current state of an asset for old_data in the audit log."""
    if asset_type == "splitter":
        return {
            "box_id": obj.box_id,
            "splitter_level": obj.splitter_level,
            "splitter_type": obj.splitter_type,
            "input_ports": obj.input_ports,
            "output_ports": obj.output_ports,
            "number_customer": obj.number_customer,
            "longitude": obj.longitude,
            "latitude": obj.latitude,
        }
    if asset_type == "cabinet":
        return {
            "cabinet_id": obj.cabinet_id,
            "capacity": obj.capacity,
            "number_tray": obj.number_tray,
            "longitude": obj.longitude,
            "latitude": obj.latitude,
        }
    if asset_type == "olt":
        return {
            "name": obj.name,
            "location_description": obj.location_description,
            "number_of_odf": obj.number_of_odf,
            "longitude": obj.longitude,
            "latitude": obj.latitude,
        }
    return {}
