# ============================================================
# OPSYN NETWORK UPLOADS — app/modules/network/upload.py
#
# Bulk import straight into net_assets / net_cables (Network GIS Phase 2),
# through the existing upload session → staging → commit flow
# (infrastructure/upload_service.py delegates here for these types):
#
#   poles, manholes, handholes, FDTs, FATs, closures, enclosures, POPs,
#   ODFs, ONUs, customer end points  → net_assets
#   cables                           → net_cables (as-built, IMPORTED)
#
# Formats: Excel (.xlsx, first sheet, header row) or GeoJSON
# (.geojson/.json FeatureCollection). Points come from longitude/latitude
# columns or a Point geometry; cables from a geometry_wkt column
# ("LINESTRING(lng lat, …)") or a LineString geometry.
#
# Columns: the base columns below plus the type's attribute fields from
# attributes.py — the same coercion rules as the API, so an upload can
# store nothing the inspector could not. Blank asset_code → generated at
# commit (codes.py). Rows default to ACTIVE: an upload records plant that
# exists. Everything is validated before anything is written; commit is
# all-or-nothing and idempotent per staging row (source_table/source_id).
# ============================================================

from __future__ import annotations

import io
import json
import math
import re
import uuid
from dataclasses import dataclass
from typing import Any, Iterator

import openpyxl
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.infrastructure.audit_service import write_audit_entry
from app.modules.network.asset_types import ASSET_TYPES, CABLE_CATEGORIES, INSTALLATION_TYPES
from app.modules.network.attributes import CABLE_ATTRIBUTES, FieldSpec, asset_specs, coerce
from app.modules.network.codes import next_codes
from app.modules.network.lifecycle import IMPORTED_DEFAULT_STATE, LIFECYCLE_STATES
from app.modules.network.models import NetAsset, NetCable

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_VERTICES = 5_000
STAGING_SOURCE = "infrastructure_upload_staging"
_XLSX_MAGIC = b"PK\x03\x04"


@dataclass(frozen=True)
class UploadType:
    key: str
    label: str
    description: str
    kind: str                 # asset | cable
    asset_type: str | None = None


NET_UPLOAD_TYPES: dict[str, UploadType] = {u.key: u for u in (
    UploadType("pop", "POPs", "Points of presence", "asset", "POP"),
    UploadType("odf", "ODFs", "Optical distribution frames", "asset", "ODF"),
    UploadType("fdt", "FDTs", "Fibre distribution terminals", "asset", "FDT"),
    UploadType("fat", "FATs", "Fibre access terminals", "asset", "FAT"),
    UploadType("closure", "Closures", "Dome and inline splice closures", "asset", "CLOSURE"),
    UploadType("enclosure", "Enclosures", "Wall boxes, terminal boxes, pedestals", "asset", "ENCLOSURE"),
    UploadType("pole", "Poles", "Owned, leased and shared poles", "asset", "POLE"),
    UploadType("manhole", "Manholes", "Underground chambers", "asset", "MANHOLE"),
    UploadType("handhole", "Handholes", "Small underground pits", "asset", "HANDHOLE"),
    UploadType("onu", "ONUs", "Customer optical network units", "asset", "ONU"),
    UploadType("customer_endpoint", "Customer end points", "Premises served by the network", "asset", "CUSTOMER_ENDPOINT"),
    UploadType("cable", "Cables", "Feeder, distribution, drop and other cables with their route", "cable"),
)}

ASSET_BASE_REQUIRED = ("longitude", "latitude")
ASSET_BASE_OPTIONAL = ("asset_code", "name", "description", "lifecycle_status", "capacity", "used_capacity")
CABLE_BASE_REQUIRED = ("cable_category", "geometry_wkt")
CABLE_BASE_OPTIONAL = ("asset_code", "name", "description", "lifecycle_status", "fiber_count",
                       "installation_type", "source_asset_code", "dest_asset_code")


def attribute_specs(u: UploadType) -> tuple[FieldSpec, ...]:
    return asset_specs(u.asset_type) if u.kind == "asset" else CABLE_ATTRIBUTES


def columns(u: UploadType) -> tuple[list[str], list[str]]:
    """(required, optional) column names for an Excel upload."""
    attrs = [f.key for f in attribute_specs(u)]
    if u.kind == "asset":
        return list(ASSET_BASE_REQUIRED), list(ASSET_BASE_OPTIONAL) + attrs
    return list(CABLE_BASE_REQUIRED), list(CABLE_BASE_OPTIONAL) + attrs


def upload_type_specs() -> list[dict[str, Any]]:
    out = []
    for u in NET_UPLOAD_TYPES.values():
        req, opt = columns(u)
        out.append({
            "key": u.key, "label": u.label, "description": u.description, "group": "network",
            "kind": u.kind, "asset_type": u.asset_type, "formats": ["xlsx", "geojson"],
            "required_columns": req, "optional_columns": opt,
            "geojson_geometry": "Point" if u.kind == "asset" else "LineString",
        })
    return out


def template_xlsx(u: UploadType) -> io.BytesIO:
    """A workbook with the header row and one example row."""
    req, opt = columns(u)
    headers = req + opt
    example: dict[str, Any] = {
        "longitude": 3.3530, "latitude": 6.6010, "lifecycle_status": "ACTIVE",
        "cable_category": "DISTRIBUTION", "fiber_count": 24,
        "geometry_wkt": "LINESTRING(3.3530 6.6010, 3.3544 6.6008, 3.3538 6.5992)",
    }
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = u.label[:31]
    ws.append(headers)
    ws.append([example.get(h, "") for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── Reading ──────────────────────────────────────────────────

class FileError(Exception):
    pass


def file_format(filename: str) -> str:
    name = filename.lower()
    if name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith(".geojson") or name.endswith(".json"):
        return "geojson"
    raise FileError(f"'{filename}' is not an Excel (.xlsx) or GeoJSON (.geojson, .json) file")


def _norm_key(k: Any) -> str:
    return str(k).strip().lower()


def read_rows(filename: str, data: bytes) -> tuple[list[str], Iterator[tuple[int, dict[str, Any], dict | None]]]:
    """(header names, iterator of (row number, row dict, GeoJSON geometry or None))."""
    if not data:
        raise FileError("The file is empty")
    if len(data) > MAX_BYTES:
        raise FileError(f"The file is {len(data) / 1048576:.1f} MB; the limit is {MAX_BYTES // 1048576} MB")
    fmt = file_format(filename)

    if fmt == "xlsx":
        if data[:4] != _XLSX_MAGIC:
            raise FileError("The file is not a valid Excel workbook")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise FileError(f"The workbook could not be read: {exc}") from exc
        ws = wb.active
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None) or ()
        header_map = {_norm_key(c): i for i, c in enumerate(header) if c is not None}

        def gen() -> Iterator[tuple[int, dict[str, Any], dict | None]]:
            for n, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if all(c is None for c in raw):
                    continue
                yield n, {k: raw[i] for k, i in header_map.items() if i < len(raw)}, None
            wb.close()
        return list(header_map), gen()

    try:
        doc = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise FileError(f"The file is not valid JSON: {exc}") from exc
    if not isinstance(doc, dict) or doc.get("type") != "FeatureCollection" or not isinstance(doc.get("features"), list):
        raise FileError("The GeoJSON must be a FeatureCollection")
    feats = doc["features"]
    keys = sorted({_norm_key(k) for f in feats if isinstance(f, dict) for k in (f.get("properties") or {})})

    def gen_geo() -> Iterator[tuple[int, dict[str, Any], dict | None]]:
        for n, f in enumerate(feats, start=1):
            props = (f.get("properties") or {}) if isinstance(f, dict) else {}
            geom = f.get("geometry") if isinstance(f, dict) else None
            yield n, {_norm_key(k): v for k, v in props.items()}, geom
    return keys, gen_geo()


# ── Geometry ─────────────────────────────────────────────────

_WKT_LINE = re.compile(r"^\s*LINESTRING\s*\((.+)\)\s*$", re.IGNORECASE | re.DOTALL)


def _valid_lnglat(x: float, y: float) -> str | None:
    if not (math.isfinite(x) and math.isfinite(y)):
        return "coordinates must be numbers"
    if not -180 <= x <= 180 or not -90 <= y <= 90:
        return f"({x}, {y}) is outside longitude −180..180 / latitude −90..90"
    if x == 0 and y == 0:
        return "(0, 0) is not a real location — the coordinates are missing"
    return None


def point_from_row(row: dict[str, Any], geom: dict | None) -> tuple[list[float] | None, str | None, str]:
    """([lng, lat], error, field)."""
    if geom is not None:
        if geom.get("type") != "Point":
            return None, f"geometry must be a Point (got {geom.get('type')})", "geometry"
        try:
            x, y = (float(v) for v in geom["coordinates"][:2])
        except (KeyError, TypeError, ValueError):
            return None, "the Point has no usable coordinates", "geometry"
    else:
        try:
            x, y = float(row.get("longitude")), float(row.get("latitude"))  # type: ignore[arg-type]
        except (TypeError, ValueError):
            return None, "longitude and latitude are required numbers", "longitude"
    err = _valid_lnglat(x, y)
    return (None, err, "geometry") if err else ([x, y], None, "geometry")


def line_from_row(row: dict[str, Any], geom: dict | None) -> tuple[list[list[float]] | None, str | None]:
    coords: list[list[float]] = []
    if geom is not None:
        if geom.get("type") != "LineString":
            return None, f"geometry must be a LineString (got {geom.get('type')})"
        try:
            coords = [[float(p[0]), float(p[1])] for p in geom["coordinates"]]
        except (KeyError, TypeError, ValueError, IndexError):
            return None, "the LineString has unusable coordinates"
    else:
        m = _WKT_LINE.match(str(row.get("geometry_wkt") or ""))
        if not m:
            return None, "geometry_wkt must look like LINESTRING(lng lat, lng lat, …)"
        try:
            for pair in m.group(1).split(","):
                x, y = pair.split()[:2]
                coords.append([float(x), float(y)])
        except ValueError:
            return None, "geometry_wkt has a vertex that is not two numbers"
    if len(coords) < 2 or all(c == coords[0] for c in coords):
        return None, "a cable needs at least two distinct vertices"
    if len(coords) > MAX_VERTICES:
        return None, f"a cable may have at most {MAX_VERTICES} vertices"
    for x, y in coords:
        err = _valid_lnglat(x, y)
        if err:
            return None, err
    return coords, None


# ── Row validation ───────────────────────────────────────────

def _opt_int(row: dict, key: str, label: str, min_val: int, errors: list[dict]) -> int | None:
    value, err = coerce(FieldSpec(key, label, "int", min=min_val), row.get(key))
    if err:
        errors.append({"field": key, "error": err})
    return value


def validate_row(u: UploadType, row: dict[str, Any], geom: dict | None) -> tuple[dict[str, Any], list[dict[str, str]]]:
    """Normalise one row. Returns (staged row data, errors)."""
    errors: list[dict[str, str]] = []
    out: dict[str, Any] = {}

    code = row.get("asset_code")
    code = str(code).strip() if code is not None else ""
    if len(code) > 100:
        errors.append({"field": "asset_code", "error": "asset_code is longer than 100 characters"})
    out["asset_code"] = code or None
    for key, limit in (("name", 200), ("description", 2000)):
        v = row.get(key)
        v = str(v).strip() if v is not None else ""
        if len(v) > limit:
            errors.append({"field": key, "error": f"{key} is longer than {limit} characters"})
        out[key] = v or None

    life = str(row.get("lifecycle_status") or "").strip().upper()
    if life and life not in LIFECYCLE_STATES:
        errors.append({"field": "lifecycle_status", "error": f"lifecycle_status '{life}' is not a lifecycle state"})
    out["lifecycle_status"] = life or IMPORTED_DEFAULT_STATE

    if u.kind == "asset":
        pt, err, field = point_from_row(row, geom)
        if err:
            errors.append({"field": field, "error": err})
        out["_geometry"] = {"type": "Point", "coordinates": pt} if pt else None
        out["capacity"] = _opt_int(row, "capacity", "capacity", 0, errors)
        out["used_capacity"] = _opt_int(row, "used_capacity", "used_capacity", 0, errors)
        if out["capacity"] is not None and out["used_capacity"] is not None and out["used_capacity"] > out["capacity"]:
            errors.append({"field": "used_capacity", "error": "used_capacity cannot be more than capacity"})
    else:
        cat = str(row.get("cable_category") or "").strip().upper()
        if cat not in CABLE_CATEGORIES:
            errors.append({"field": "cable_category",
                           "error": f"cable_category must be one of {', '.join(CABLE_CATEGORIES)}"})
        out["cable_category"] = cat
        line, err = line_from_row(row, geom)
        if err:
            errors.append({"field": "geometry", "error": err})
        out["_geometry"] = {"type": "LineString", "coordinates": line} if line else None
        out["fiber_count"] = _opt_int(row, "fiber_count", "fiber_count", 1, errors)
        inst = str(row.get("installation_type") or "").strip().upper()
        if inst and inst not in INSTALLATION_TYPES:
            errors.append({"field": "installation_type",
                           "error": f"installation_type must be one of {', '.join(INSTALLATION_TYPES)}"})
        out["installation_type"] = inst or None
        for key in ("source_asset_code", "dest_asset_code"):
            v = row.get(key)
            out[key] = str(v).strip() if v not in (None, "") else None

    attrs: dict[str, Any] = {}
    for spec in attribute_specs(u):
        value, err = coerce(spec, row.get(spec.key))
        if err:
            errors.append({"field": spec.key, "error": err})
        elif value is not None:
            attrs[spec.key] = value
    out["attributes"] = attrs
    return out, errors


# ── Database checks ──────────────────────────────────────────

async def existing_codes(db: AsyncSession, tid: uuid.UUID, u: UploadType, codes: list[str]) -> set[str]:
    """Lower-cased codes already used by a live asset of this type (or any live cable)."""
    if not codes:
        return set()
    if u.kind == "asset":
        sql = """SELECT lower(asset_code) FROM net_assets
                 WHERE tenant_id = :tid AND asset_type = :t AND design_id IS NULL AND deleted_at IS NULL
                   AND lower(asset_code) = ANY(:codes)"""
        params = {"tid": tid, "t": u.asset_type, "codes": codes}
    else:
        sql = """SELECT lower(asset_code) FROM net_cables
                 WHERE tenant_id = :tid AND design_id IS NULL AND deleted_at IS NULL
                   AND lower(asset_code) = ANY(:codes)"""
        params = {"tid": tid, "codes": codes}
    return {r[0] for r in (await db.execute(text(sql), params)).all()}


async def resolve_asset_codes(db: AsyncSession, tid: uuid.UUID, codes: list[str]) -> dict[str, list[str]]:
    """lower(code) → ids of live assets with that code (any type)."""
    if not codes:
        return {}
    rows = (await db.execute(text("""
        SELECT lower(asset_code) AS code, id::text AS id FROM net_assets
        WHERE tenant_id = :tid AND design_id IS NULL AND deleted_at IS NULL AND lower(asset_code) = ANY(:codes)
    """), {"tid": tid, "codes": codes})).all()
    out: dict[str, list[str]] = {}
    for r in rows:
        out.setdefault(r.code, []).append(r.id)
    return out


# ── Validation (writes staging rows) ─────────────────────────

async def validate_upload(db: AsyncSession, session_obj: Any, u: UploadType, filename: str,
                          data: bytes, tid: uuid.UUID) -> Any:
    """Validate a whole file into staging. Returns infrastructure ValidationSummary."""
    import hashlib

    from sqlalchemy import delete as sa_delete

    from app.modules.infrastructure.models import InfraUploadStaging
    from app.modules.infrastructure.schemas import ValidationSummary

    def failed(msg: str, field: str = "file", total: int = 0) -> ValidationSummary:
        session_obj.status = "failed"
        session_obj.validation_errors = [{"error": msg}]
        return ValidationSummary(session_id=session_obj.id, asset_type=u.key, status="failed",
                                 total_rows=total, valid_rows=0, duplicate_rows=0, error_rows=0,
                                 warnings=0, errors=[{"row": 0, "field": field, "error": msg}], duplicates=[])

    session_obj.status = "validating"
    session_obj.file_hash = hashlib.sha256(data).hexdigest()
    try:
        headers, rows_iter = read_rows(filename, data)
    except FileError as exc:
        return failed(str(exc))

    if file_format(filename) == "xlsx":
        required, _opt = columns(u)
        missing = [c for c in required if c not in headers]
        if missing:
            return failed(f"Missing required columns: {', '.join(missing)}", "header")

    parsed: list[tuple[int, dict[str, Any], list[dict[str, str]]]] = []
    for n, row, geom in rows_iter:
        if len(parsed) >= MAX_ROWS:
            return failed(f"The file has more than {MAX_ROWS} rows; split it into smaller files", total=MAX_ROWS + 1)
        out, errs = validate_row(u, row, geom)
        parsed.append((n, out, errs))

    # The same code twice in one file: the second one is an error
    seen: dict[str, int] = {}
    for n, out, errs in parsed:
        code = (out.get("asset_code") or "").lower()
        if not code:
            continue
        if code in seen:
            errs.append({"field": "asset_code", "error": f"asset_code '{out['asset_code']}' is also used on row {seen[code]}"})
        else:
            seen[code] = n

    taken = await existing_codes(db, tid, u, list(seen))

    if u.kind == "cable":
        wanted = sorted({(out.get(k) or "").lower() for _n, out, _e in parsed
                         for k in ("source_asset_code", "dest_asset_code") if out.get(k)})
        resolved = await resolve_asset_codes(db, tid, wanted)
        for _n, out, errs in parsed:
            for key in ("source_asset_code", "dest_asset_code"):
                code = out.get(key)
                if not code:
                    continue
                hits = resolved.get(code.lower(), [])
                if not hits:
                    errs.append({"field": key, "error": f"no asset in the live network has code '{code}'"})
                elif len(hits) > 1:
                    errs.append({"field": key, "error": f"{len(hits)} live assets share code '{code}'; link it by hand"})

    staging: list[Any] = []
    errors: list[dict[str, Any]] = []
    duplicates: list[dict[str, Any]] = []
    valid = dup = bad = 0
    for n, out, errs in parsed:
        is_dup = not errs and (out.get("asset_code") or "").lower() in taken
        if errs:
            bad += 1
            errors.extend({"row": n, **e} for e in errs)
        elif is_dup:
            dup += 1
            duplicates.append({"row": n, "asset_code": out["asset_code"]})
        else:
            valid += 1
        staging.append(InfraUploadStaging(
            session_id=session_obj.id, tenant_id=tid, asset_type=u.key, row_number=n,
            row_data=out, is_valid=not errs, is_duplicate=is_dup, errors=errs, committed=False,
        ))

    await db.execute(sa_delete(InfraUploadStaging).where(InfraUploadStaging.session_id == session_obj.id))
    db.add_all(staging)
    session_obj.status = "validated"
    session_obj.total_rows = len(parsed)
    session_obj.valid_rows = valid
    session_obj.duplicate_rows = dup
    session_obj.error_rows = bad
    session_obj.validation_errors = errors
    await db.flush()
    return ValidationSummary(session_id=session_obj.id, asset_type=u.key, status="validated",
                             total_rows=len(parsed), valid_rows=valid, duplicate_rows=dup, error_rows=bad,
                             warnings=0, errors=errors, duplicates=duplicates)


# ── Commit ───────────────────────────────────────────────────

def _ewkt(geom: dict[str, Any]) -> str:
    if geom["type"] == "Point":
        x, y = geom["coordinates"]
        return f"SRID=4326;POINT({x} {y})"
    return "SRID=4326;LINESTRING(" + ", ".join(f"{x} {y}" for x, y in geom["coordinates"]) + ")"


async def commit_rows(db: AsyncSession, u: UploadType, staging_rows: list[Any], tid: uuid.UUID,
                      actor: Any, session_id: uuid.UUID) -> int:
    """Insert every staged row; caller owns the savepoint. Returns rows written."""
    rows = [(s, s.row_data) for s in staging_rows]
    blanks = [d for _s, d in rows if not d.get("asset_code")]
    if u.kind == "asset":
        generated = await next_codes(db, tid, "asset", u.asset_type, len(blanks)) if blanks else []
        for d, code in zip(blanks, generated):
            d["_generated_code"] = code
    else:
        by_cat: dict[str, list[dict]] = {}
        for d in blanks:
            by_cat.setdefault(d["cable_category"], []).append(d)
        for cat, ds in by_cat.items():
            for d, code in zip(ds, await next_codes(db, tid, "cable", cat, len(ds))):
                d["_generated_code"] = code
        endpoint_codes = sorted({(d.get(k) or "").lower() for _s, d in rows
                                 for k in ("source_asset_code", "dest_asset_code") if d.get(k)})
        resolved = await resolve_asset_codes(db, tid, endpoint_codes)

    written = 0
    for staging, d in rows:
        code = d.get("asset_code") or d["_generated_code"]
        common = dict(
            tenant_id=tid, asset_code=code, name=d.get("name"), description=d.get("description"),
            geom=_ewkt(d["_geometry"]), lifecycle_status=d["lifecycle_status"], creation_method="IMPORTED",
            attributes=d.get("attributes") or {}, source_table=STAGING_SOURCE, source_id=staging.id,
            created_by=actor.id if actor else None, updated_by=actor.id if actor else None,
        )
        if u.kind == "asset":
            obj: Any = NetAsset(asset_type=u.asset_type, capacity=d.get("capacity"),
                                used_capacity=d.get("used_capacity"), **common)
        else:
            def one(key: str) -> uuid.UUID | None:
                ids = resolved.get((d.get(key) or "").lower(), [])
                return uuid.UUID(ids[0]) if len(ids) == 1 else None
            obj = NetCable(cable_category=d["cable_category"], fiber_count=d.get("fiber_count"),
                           installation_type=d.get("installation_type"), route_method="IMPORTED",
                           source_asset_id=one("source_asset_code"), dest_asset_id=one("dest_asset_code"),
                           **common)
        db.add(obj)
        await db.flush()
        await write_audit_entry(
            db, asset_type=f"net_{u.kind}:{u.asset_type or d['cable_category']}", asset_id=obj.id,
            asset_key=code, action="upload", session_id=session_id, actor_user=actor,
            old_data=None, new_data={k: v for k, v in d.items() if k != "_geometry"}, tenant_id=tid,
        )
        written += 1
    from app.modules.network.topology import sync_derived_connections
    await sync_derived_connections(db, tid)
    return written


assert all(u.asset_type in ASSET_TYPES for u in NET_UPLOAD_TYPES.values() if u.kind == "asset")
